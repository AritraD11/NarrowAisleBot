"""Type, colour and the handful of layouts the deck is built from.

Light ground, one navy accent, serif headings. The four status colours are the
report's own: green for something measured, amber for something bounded or
still open, red for a defect or a gap, grey for something configured but not
exercised. Keeping them the same across the report and the slides means a
colour on a slide means what it means in the figure next to it.
"""
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
import copy

# --- colour -----------------------------------------------------------------
INK      = RGBColor(0x1A, 0x22, 0x30)
BODY     = RGBColor(0x3B, 0x45, 0x57)
MUTED    = RGBColor(0x83, 0x8E, 0x9F)
ACCENT   = RGBColor(0x1F, 0x4E, 0x79)
ACCENT_2 = RGBColor(0x4C, 0x7A, 0xA8)
TINT     = RGBColor(0xF3, 0xF6, 0xFA)
TINT_2   = RGBColor(0xE7, 0xED, 0xF4)
RULE     = RGBColor(0xD6, 0xDD, 0xE6)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)

GREEN    = RGBColor(0x2E, 0x7D, 0x52)
AMBER    = RGBColor(0xA9, 0x6F, 0x15)
RED      = RGBColor(0xB0, 0x3A, 0x34)
GREY     = RGBColor(0x7A, 0x85, 0x96)

STATUS = {'measured': GREEN, 'open': AMBER, 'gap': RED, 'idle': GREY, 'accent': ACCENT}

# --- type -------------------------------------------------------------------
HEAD = 'Cambria'
TEXT = 'Calibri'

# --- geometry (inches) ------------------------------------------------------
SW, SH   = 13.333, 7.5
MARGIN   = 0.62
TITLE_Y  = 0.42
RULE_Y   = 1.14
BODY_Y   = 1.44
BODY_B   = 6.62          # bottom of the content region
FOOT_Y   = 6.90
COL      = SW - 2 * MARGIN


def _in(v):
    return Inches(v)


def box(slide, x, y, w, h, fill=None, line=None, lw=0.75):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _in(x), _in(y), _in(w), _in(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.text_frame.word_wrap = True
    return sh


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(_in(x), _in(y), _in(w), _in(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def run(p, text, size=14, font=TEXT, color=BODY, bold=False, italic=False, spacing=None):
    r = p.add_run(); r.text = text
    f = r.font
    f.name = font; f.size = Pt(size); f.color.rgb = color
    f.bold = bold; f.italic = italic
    if spacing is not None:                      # character spacing, in points
        r._r.get_or_add_rPr().set('spc', str(int(spacing * 100)))
    return r


def para(tf, first=False, space_before=0, space_after=0, line=1.0, align=PP_ALIGN.LEFT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    p.alignment = align
    return p


def bullet(p, color=ACCENT, char='•', indent=0.20):
    """Give a paragraph a real PowerPoint bullet, coloured and hung."""
    pPr = p._p.get_or_add_pPr()
    pPr.set('marL', str(Emu(int(indent * 914400))))
    pPr.set('indent', str(-Emu(int(indent * 914400))))
    for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum'):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    clr = pPr.makeelement(qn('a:buClr'), {})
    srgb = pPr.makeelement(qn('a:srgbClr'), {'val': str(color)})
    clr.append(srgb)
    fnt = pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'})
    bu = pPr.makeelement(qn('a:buChar'), {'char': char})
    for el in (clr, fnt, bu):
        pPr.append(el)
    return p


def fit_image(slide, path, x, y, w, h, frame=False):
    """Place an image inside the box, preserving aspect, centred."""
    from PIL import Image
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    dw, dh = iw * scale, ih * scale
    dx, dy = x + (w - dw) / 2, y + (h - dh) / 2
    if frame:
        box(slide, dx - 0.045, dy - 0.045, dw + 0.09, dh + 0.09, fill=WHITE, line=RULE, lw=0.75)
    return slide.shapes.add_picture(path, _in(dx), _in(dy), _in(dw), _in(dh))


def cover_image(slide, path, x, y, w, h):
    """Fill the box completely, cropping the overflow (PowerPoint-side crop)."""
    from PIL import Image
    iw, ih = Image.open(path).size
    scale = max(w / iw, h / ih)
    dw, dh = iw * scale, ih * scale
    pic = slide.shapes.add_picture(path, _in(x - (dw - w) / 2), _in(y - (dh - h) / 2), _in(dw), _in(dh))
    # crop back to the box so nothing spills off-slide
    cl = max(0.0, (dw - w) / 2 / dw); ct = max(0.0, (dh - h) / 2 / dh)
    pic.crop_left = cl; pic.crop_right = cl
    pic.crop_top = ct;  pic.crop_bottom = ct
    pic.left, pic.top, pic.width, pic.height = _in(x), _in(y), _in(w), _in(h)
    return pic


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text.strip()
