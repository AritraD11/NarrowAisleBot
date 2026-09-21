#!/usr/bin/env python3
"""Render the deck through LibreOffice and check every slide for overflow.

    python3 docs/aps_report/seminar_final/src/qa.py

Three checks, and the first is the one that matters:

  1. A measurement pass over src/content.py, before anything is written: every
     bullet block is wrapped against the box it will be given, using fonts
     metric-compatible with Cambria and Calibri, and flagged if it does not
     fit. python-pptx cannot autofit, so a paragraph one line too long does not
     error - it silently spills over whatever is beneath it.
  2. A geometric pass over the .pptx: any shape whose box leaves the slide.
  3. A pixel pass over the rendered PNGs: ink below the footer line or inside
     the right margin, which is what spilled text looks like once drawn.

The PNGs are written to build/png/ and are worth looking at, not just grepping.
LibreOffice's text metrics are not PowerPoint's, so treat a marginal overflow
here as a warning and a large one as a fault.
"""
import glob
import os
import shutil
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..'))
DECK = os.path.join(ROOT, 'NarrowAisleBot_APS_Seminar.pptx')
BUILD = os.path.join(ROOT, 'build')
PNG = os.path.join(BUILD, 'png')

SW_EMU, SH_EMU = 12192000, 6858000          # 13.333 x 7.5 in
EMU_IN = 914400


def overflow():
    """Measure each text block against the box it was given.

    This is the check that matters. python-pptx will happily write a bullet
    list two inches taller than its box, and nothing in the file says so - the
    text simply runs over whatever is under it. Measuring the wrap with the
    metric-compatible fonts catches it before the render does.
    """
    sys.path.insert(0, HERE)
    import metrics
    import layouts
    from content import SLIDES

    faults = []
    for i, s in enumerate(SLIDES, start=1):
        if s.get('layout') in ('title', 'divider', 'closing'):
            continue
        items = s.get('bullets')
        if not items:
            continue
        top = layouts.RULE_Y + (len(layouts.title_lines(s['title'])) - 1) * 0.46 + 0.30
        if s.get('tiles'):
            top += 1.16
        bottom = 5.86 if s.get('takeaway') else layouts.BODY_B
        if s.get('videos'):
            top += min(2.55, bottom - top - 1.6) + 0.24
        if s.get('images'):
            top += min(3.1, bottom - top - 1.9) + 0.36
        img, side = (s.get('image') or s.get('video')), s.get('side', 'right')
        if img and side in ('right', 'left') and not (s.get('videos') or s.get('images')):
            w = layouts.COL * (1 - s.get('image_w', 0.50)) - 0.42
        else:
            w = layouts.COL * s.get('text_w', 0.88)
        size = s.get('size', 14.5 if (img and side in ('right', 'left')) else 15.5)
        gap = s.get('gap', 9 if (img and side in ('right', 'left')) else 12)

        used = 0.0
        for it in items:
            text = it if isinstance(it, str) else it[0]
            lead = '' if isinstance(it, str) or len(it) < 3 or not it[2] else it[2]
            sub = text.startswith('~')
            body = (lead + text[1:]) if sub else (lead + text)
            used += metrics.block_height_in(
                body, w - (0.20 if not sub else 0.20),
                size=size - 1.5 if sub else size, name='Calibri',
                spacing=1.05, space_after_pt=gap)
        avail = bottom - top
        if used > avail + 0.02:
            faults.append('slide %2d  bullets need %.2f in, box is %.2f in  (%s)'
                          % (i, used, avail, s['title'][:46]))

        take = s.get('takeaway')
        if take:
            lines = metrics.wrap(take, layouts.COL - 0.50, 'Cambria', 13, italic=True)
            if len(lines) > 2:
                faults.append('slide %2d  takeaway wraps to %d lines, box holds 2  (%s)'
                              % (i, len(lines), s['title'][:46]))
    return faults


def geometry():
    from pptx import Presentation
    prs = Presentation(DECK)
    faults = []
    for i, slide in enumerate(prs.slides, start=1):
        for sh in slide.shapes:
            if sh.left is None:
                continue
            r, b = sh.left + sh.width, sh.top + sh.height
            if sh.left < -1000 or sh.top < -1000 or r > SW_EMU + 1000 or b > SH_EMU + 1000:
                faults.append('slide %2d  shape off-slide: %s  (%.2f, %.2f) %.2f x %.2f in'
                              % (i, sh.shape_type, sh.left / EMU_IN, sh.top / EMU_IN,
                                 sh.width / EMU_IN, sh.height / EMU_IN))
    return faults


def render():
    if shutil.which('soffice') is None:
        print('  soffice not found - skipping the render pass')
        return []
    os.makedirs(BUILD, exist_ok=True)
    shutil.rmtree(PNG, ignore_errors=True)
    os.makedirs(PNG)
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', BUILD, DECK],
                   check=True, capture_output=True, timeout=600)
    pdf = os.path.join(BUILD, os.path.basename(DECK).replace('.pptx', '.pdf'))
    import pymupdf
    doc = pymupdf.open(pdf)
    for i, page in enumerate(doc, start=1):
        page.get_pixmap(dpi=110).save(os.path.join(PNG, 'slide%02d.png' % i))
    return sorted(glob.glob(os.path.join(PNG, '*.png')))


def ink_bounds(pages, full_bleed=()):
    """Flag ink in the margins the layout reserves.

    The title, divider and closing slides are full-bleed by design, so they are
    passed in as exceptions rather than reported every run.
    """
    import numpy as np
    from PIL import Image
    faults = []
    for p in pages:
        n = int(os.path.basename(p)[5:7])
        a = np.asarray(Image.open(p).convert('L'))
        h, w = a.shape
        ink = a < 235
        # the dividers, title and closing slides are full-bleed by design
        if n in set(full_bleed):
            continue
        right = ink[:, int(w * 0.975):]
        bot = ink[int(h * 0.975):, :]
        if right.mean() > 0.02:
            faults.append('slide %2d  ink in the right margin (%.1f %%)' % (n, right.mean() * 100))
        if bot.mean() > 0.02:
            faults.append('slide %2d  ink below the footer line (%.1f %%)' % (n, bot.mean() * 100))
    return faults


def main():
    faults = overflow() + geometry()
    pages = render()
    if pages:
        sys.path.insert(0, HERE)
        from content import SLIDES
        bleed = [i for i, sl in enumerate(SLIDES, start=1)
                 if sl.get('layout') in ('title', 'divider', 'closing')]
        faults += ink_bounds(pages, bleed)
    print('%d slides rendered to %s' % (len(pages), os.path.relpath(PNG, os.getcwd())))
    if faults:
        print('\n%d issue(s):' % len(faults))
        for f in faults:
            print('  ' + f)
        return 1
    print('no geometric or margin faults')
    return 0


if __name__ == '__main__':
    sys.exit(main())
